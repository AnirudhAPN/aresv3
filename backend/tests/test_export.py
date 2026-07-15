"""
Tests for services/export.py, focused on the Word (.docx) export: it must
produce a real, valid .docx file — not just avoid throwing — and every
markdown/text/PDF export must degrade gracefully on missing/empty fields
rather than crashing (a completed session can legitimately have empty
lists for any section).
"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from services import export as export_service


def _full_session() -> dict:
    return {
        "topic": "Explainable AI for Chest X-Ray Diagnosis",
        "created_at": 1784134739.0,
        "parse_failures": {},
        "agent_outputs": {
            "planner": {"output": {
                "research_objectives": ["Investigate interpretability methods"],
                "research_questions": ["Which XAI method balances fidelity and trust?"],
                "research_scope": "Focused on Grad-CAM and SHAP-based explanations.",
            }},
            "literature": {"output": {
                "papers": [{
                    "title": "Explainable Deep Learning for Chest X-Ray Diagnosis",
                    "authors": ["A. Author"], "year": 2022, "url": "https://example.com/a",
                    "summary": "Grad-CAM based approach.",
                    "citation_apa": "Author, A. (2022). Explainable Deep Learning.",
                    "url_verified": True,
                }],
                "low_coverage": False,
                "similar_work": {"paper_title": "Similar Paper", "paper_url": "https://example.com/sim", "similarity": 0.83},
            }},
            "gap": {"output": {
                "current_trends": ["Growing use of Grad-CAM"],
                "gaps": [{"gap_description": "Not addressed: patient-specific calibration.",
                           "based_on_papers": ["Paper A"], "confidence_note": "based on 1 paper"}],
                "future_directions": ["Clinician-in-the-loop tuning"], "low_confidence_gaps": False,
            }},
            "methodology": {"output": {
                "datasets": [{"name": "PhysioNet", "url": "https://physionet.org/", "verified": True,
                               "url_verified": True, "kaggle_url": "https://www.kaggle.com/datasets/mondejar/mitbih-database"}],
                "algorithms": [{"name": "Grad-CAM", "url": "https://arxiv.org/abs/1610.02391",
                                  "verified": True, "url_verified": True}],
                "evaluation_metrics": ["AUROC", "clinician agreement"],
                "implementation_steps": ["Train baseline CNN", "Apply Grad-CAM", "Evaluate against annotations"],
                "unverified_suggestions": [],
            }},
            "reviewer": {"output": {
                "novelty_score": 0.72, "quality_score": 0.68,
                "rule_check_results": {"gaps_have_evidence": True, "methodology_ok": False},
                "strengths": ["Clear framing"], "weaknesses": ["Small literature base"], "missing_components": [],
            }},
        },
    }


def test_docx_export_is_a_valid_zip_with_expected_content():
    from docx import Document

    docx_bytes = export_service.render_docx_bytes(_full_session())
    assert docx_bytes[:2] == b"PK"  # docx is a zip container

    doc = Document(io.BytesIO(docx_bytes))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Research Proposal: Explainable AI for Chest X-Ray Diagnosis" in all_text
    assert "Investigate interpretability methods" in all_text
    assert "Explainable Deep Learning for Chest X-Ray Diagnosis" in all_text
    assert "Citation (APA)" in all_text
    assert "Train baseline CNN" in all_text
    assert "PASS" in all_text and "FAIL" in all_text  # rule check results rendered


def test_docx_export_creates_real_hyperlinks_not_plain_text_urls():
    from docx import Document

    docx_bytes = export_service.render_docx_bytes(_full_session())
    doc = Document(io.BytesIO(docx_bytes))
    hyperlink_rels = [r for r in doc.part.rels.values() if "hyperlink" in r.reltype]
    # paper url, similar-work url, dataset kaggle url, algorithm url
    assert len(hyperlink_rels) == 4
    targets = {r.target_ref for r in hyperlink_rels}
    assert "https://example.com/a" in targets
    assert "https://www.kaggle.com/datasets/mondejar/mitbih-database" in targets


def test_docx_export_handles_missing_urls_without_crashing():
    from docx import Document

    session = _full_session()
    session["agent_outputs"]["methodology"]["output"]["datasets"] = [
        {"name": "Unnamed Dataset", "url": None, "verified": False, "url_verified": None, "kaggle_url": None}
    ]
    session["agent_outputs"]["methodology"]["output"]["algorithms"] = [
        {"name": "Custom Heuristic", "url": None, "verified": False, "url_verified": None}
    ]

    docx_bytes = export_service.render_docx_bytes(session)
    doc = Document(io.BytesIO(docx_bytes))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Unnamed Dataset" in all_text
    assert "Custom Heuristic" in all_text


def test_docx_export_handles_empty_session_without_crashing():
    from docx import Document

    empty_session = {"topic": "Untitled", "created_at": None, "parse_failures": {}, "agent_outputs": {}}
    docx_bytes = export_service.render_docx_bytes(empty_session)
    assert docx_bytes[:2] == b"PK"
    Document(io.BytesIO(docx_bytes))  # re-parses without raising


def test_markdown_and_text_exports_still_work_alongside_docx():
    session = _full_session()
    md = export_service.to_markdown(session)
    text = export_service.to_text(session)
    assert "Explainable Deep Learning for Chest X-Ray Diagnosis" in md
    assert "[" not in text and "](" not in text  # markdown link syntax stripped in the text export


if __name__ == "__main__":
    test_docx_export_is_a_valid_zip_with_expected_content()
    test_docx_export_creates_real_hyperlinks_not_plain_text_urls()
    test_docx_export_handles_missing_urls_without_crashing()
    test_docx_export_handles_empty_session_without_crashing()
    test_markdown_and_text_exports_still_work_alongside_docx()
    print("All export tests passed.")
