"""
Export a completed ResearchSession to Markdown, plain text, PDF, or Word
(.docx).

Every exported artifact carries the same caveats shown in the UI:
low_coverage / low_confidence_gaps / parse_failed flags are rendered in a
footer so exported documents can't silently drop the warnings.
"""
from __future__ import annotations

from datetime import datetime, timezone
from typing import Any


def _flags_footer(session: dict[str, Any]) -> list[str]:
    footer = []
    lit = (session.get("agent_outputs", {}).get("literature") or {}).get("output") or {}
    gap = (session.get("agent_outputs", {}).get("gap") or {}).get("output") or {}
    parse_failures = session.get("parse_failures", {})

    if lit.get("low_coverage"):
        footer.append("Literature coverage was LOW (fewer than 3 relevant papers retrieved).")
    if gap.get("low_confidence_gaps"):
        footer.append("Gap analysis is based on limited retrieved literature — treat with caution.")
    for agent, failed in parse_failures.items():
        if failed:
            footer.append(f"The '{agent}' agent's output failed to parse correctly and used a fallback.")
    return footer


def to_markdown(session: dict[str, Any]) -> str:
    topic = session.get("topic", "Untitled Research")
    created = session.get("created_at")
    date_str = datetime.fromtimestamp(created, tz=timezone.utc).strftime("%Y-%m-%d") if created else "unknown"

    outputs = session.get("agent_outputs", {})
    planner = (outputs.get("planner") or {}).get("output") or {}
    literature = (outputs.get("literature") or {}).get("output") or {}
    gap = (outputs.get("gap") or {}).get("output") or {}
    methodology = (outputs.get("methodology") or {}).get("output") or {}
    reviewer = (outputs.get("reviewer") or {}).get("output") or {}

    lines = [f"# Research Proposal: {topic}", f"_Generated {date_str} by ARES_", ""]

    lines.append("## Research Objectives")
    lines += [f"- {o}" for o in planner.get("research_objectives", [])]
    lines.append("")
    lines.append("## Research Questions")
    lines += [f"- {q}" for q in planner.get("research_questions", [])]
    lines.append("")
    lines.append(f"**Scope:** {planner.get('research_scope', '')}")
    lines.append("")

    lines.append("## Literature Review")
    if literature.get("low_coverage"):
        lines.append("> ⚠️ Low coverage: fewer than 3 relevant papers were retrieved.")
    similar_work = literature.get("similar_work")
    if similar_work:
        pct = round(similar_work.get("similarity", 0) * 100)
        lines.append(
            f"> ⚠️ Possible existing work: [{similar_work.get('paper_title')}]"
            f"({similar_work.get('paper_url')}) is {pct}% similar to this topic — check it isn't already covered."
        )
    for p in literature.get("papers", []):
        authors = ", ".join(p.get("authors", []))
        link_status = "link unreachable" if p.get("url_verified") is False else None
        lines.append(f"- **[{p.get('title')}]({p.get('url')})** ({p.get('year', 'n.d.')}) — {authors}")
        lines.append(f"  {p.get('summary', '')}")
        if p.get("citation_apa"):
            lines.append(f"  *Citation (APA):* {p['citation_apa']}")
        if link_status:
            lines.append(f"  ⚠️ {link_status}")
    lines.append("")

    lines.append("## Research Gaps")
    for t in gap.get("current_trends", []):
        lines.append(f"- Trend: {t}")
    for g in gap.get("gaps", []):
        based_on = ", ".join(g.get("based_on_papers", []))
        lines.append(f"- **Gap:** {g.get('gap_description')} _(based on: {based_on}; {g.get('confidence_note', '')})_")
    for d in gap.get("future_directions", []):
        lines.append(f"- Future direction: {d}")
    lines.append("")

    lines.append("## Methodology")
    lines.append("**Datasets:**")
    for d in methodology.get("datasets", []):
        badge = "✅ verified" if d.get("verified") else "⚠️ unverified"
        if d.get("url_verified") is False:
            badge += ", ⚠️ link unreachable"
        link = d.get("kaggle_url") or d.get("url", "")
        lines.append(f"- [{d.get('name')}]({link}) — {badge}")
    lines.append("**Algorithms:**")
    for a in methodology.get("algorithms", []):
        badge = "✅ verified" if a.get("verified") else "⚠️ unverified"
        if a.get("url_verified") is False:
            badge += ", ⚠️ link unreachable"
        lines.append(f"- {a.get('name')} — {badge}")
    lines.append("**Evaluation metrics:** " + ", ".join(methodology.get("evaluation_metrics", [])))
    lines.append("**Implementation steps:**")
    for i, step in enumerate(methodology.get("implementation_steps", []), start=1):
        lines.append(f"{i}. {step}")
    if methodology.get("unverified_suggestions"):
        lines.append("**Unverified suggestions (not checked against a real source):**")
        lines += [f"- {s}" for s in methodology["unverified_suggestions"]]
    lines.append("")

    lines.append("## Reviewer Feedback")
    lines.append(f"**Novelty score:** {reviewer.get('novelty_score', 'n/a')} | **Quality score:** {reviewer.get('quality_score', 'n/a')}")
    lines.append("**Rule checks:**")
    for check, passed in reviewer.get("rule_check_results", {}).items():
        lines.append(f"- {'✅' if passed else '❌'} {check}")
    lines.append("**Strengths:**")
    lines += [f"- {s}" for s in reviewer.get("strengths", [])]
    lines.append("**Weaknesses:**")
    lines += [f"- {w}" for w in reviewer.get("weaknesses", [])]
    lines.append("**Missing components:**")
    lines += [f"- {m}" for m in reviewer.get("missing_components", [])]
    lines.append("")

    footer = _flags_footer(session)
    if footer:
        lines.append("---")
        lines.append("### Caveats")
        lines += [f"- ⚠️ {f}" for f in footer]

    return "\n".join(lines)


def to_text(session: dict[str, Any]) -> str:
    """Same content as Markdown, but with links rendered inline as
    'Title (url)' instead of markdown link syntax."""
    md = to_markdown(session)
    import re

    def replace_link(match: "re.Match") -> str:
        text, url = match.group(1), match.group(2)
        return f"{text} ({url})" if url else text

    text = re.sub(r"\[([^\]]+)\]\(([^)]*)\)", replace_link, md)
    text = re.sub(r"[#*_>`]", "", text)
    return text


def to_pdf_html(session: dict[str, Any]) -> str:
    """Render an HTML document suitable for a WeasyPrint HTML->PDF pipeline
    (proper layout, not a raw text dump), including header/footer caveats."""
    topic = session.get("topic", "Untitled Research")
    created = session.get("created_at")
    date_str = datetime.fromtimestamp(created, tz=timezone.utc).strftime("%Y-%m-%d") if created else "unknown"
    footer = _flags_footer(session)

    import markdown2

    body_html = markdown2.markdown(to_markdown(session).split("\n", 2)[-1])  # skip the H1/subtitle, styled separately below

    footer_html = ""
    if footer:
        items = "".join(f"<li>{f}</li>" for f in footer)
        footer_html = f"""
        <div class="caveats">
          <h3>Caveats</h3>
          <ul>{items}</ul>
        </div>
        """

    return f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  body {{ font-family: 'Helvetica Neue', Arial, sans-serif; color: #1a1a2e; margin: 2cm; }}
  h1 {{ color: #4b2ea8; border-bottom: 2px solid #4b2ea8; padding-bottom: 8px; }}
  h2 {{ color: #6c3ec9; margin-top: 24px; }}
  .header {{ display: flex; justify-content: space-between; font-size: 11px; color: #666; margin-bottom: 16px; }}
  .caveats {{ background: #fff3cd; border: 1px solid #ffe69c; padding: 12px 16px; margin-top: 24px; border-radius: 6px; }}
  a {{ color: #4b2ea8; }}
  footer {{ font-size: 10px; color: #999; margin-top: 32px; border-top: 1px solid #ddd; padding-top: 8px; }}
</style>
</head>
<body>
  <div class="header">
    <span>ARES Research Proposal</span>
    <span>Generated {date_str}</span>
  </div>
  <h1>{topic}</h1>
  {body_html}
  {footer_html}
  <footer>Generated by ARES (Adaptive Research Ecosystem Suite). Verify all cited papers via their linked sources.</footer>
</body>
</html>
"""


def render_pdf_bytes(session: dict[str, Any]) -> bytes:
    """Render the session to PDF bytes using WeasyPrint. Import is local so
    the rest of the app doesn't hard-depend on WeasyPrint's native libs at
    import time."""
    from weasyprint import HTML

    html = to_pdf_html(session)
    return HTML(string=html).write_pdf()


def _add_hyperlink(paragraph, url: str, text: str, bold: bool = False) -> None:
    """Insert a real, clickable hyperlink run into a python-docx paragraph.

    python-docx has no built-in hyperlink API — this is the standard
    low-level OOXML recipe (add a relationship, build a <w:hyperlink> run
    referencing it). Falls back to plain text if url is empty, since a
    hyperlink with no target is worse than no link at all."""
    if not url:
        paragraph.add_run(text)
        return

    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    part = paragraph.part
    r_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run_el = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "4B2EA8")
    run_props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_props.append(underline)
    if bold:
        run_props.append(OxmlElement("w:b"))
    run_el.append(run_props)

    text_el = OxmlElement("w:t")
    text_el.text = text
    run_el.append(text_el)
    hyperlink.append(run_el)
    paragraph._p.append(hyperlink)


def render_docx_bytes(session: dict[str, Any]) -> bytes:
    """Render the session to a real .docx file (python-docx), with proper
    Word semantics — heading styles, bullet/numbered lists, bold labels,
    and clickable hyperlinks — built directly from the structured session
    data rather than converted from the Markdown export, so link targets
    and list ordering stay authoritative in one place per field."""
    import io

    from docx import Document
    from docx.shared import Pt, RGBColor

    topic = session.get("topic", "Untitled Research")
    created = session.get("created_at")
    date_str = datetime.fromtimestamp(created, tz=timezone.utc).strftime("%Y-%m-%d") if created else "unknown"

    outputs = session.get("agent_outputs", {})
    planner = (outputs.get("planner") or {}).get("output") or {}
    literature = (outputs.get("literature") or {}).get("output") or {}
    gap = (outputs.get("gap") or {}).get("output") or {}
    methodology = (outputs.get("methodology") or {}).get("output") or {}
    reviewer = (outputs.get("reviewer") or {}).get("output") or {}

    doc = Document()

    title = doc.add_heading(f"Research Proposal: {topic}", level=0)
    title.runs[0].font.color.rgb = RGBColor(0x4B, 0x2E, 0xA8)
    meta = doc.add_paragraph()
    meta.add_run(f"Generated {date_str} by ARES").italic = True

    def heading(text: str) -> None:
        doc.add_heading(text, level=1)

    def bullet(text: str) -> None:
        doc.add_paragraph(text, style="List Bullet")

    def bold_line(label: str, value: str) -> None:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)

    heading("Research Objectives")
    for o in planner.get("research_objectives", []):
        bullet(o)

    heading("Research Questions")
    for q in planner.get("research_questions", []):
        bullet(q)
    if planner.get("research_scope"):
        bold_line("Scope", planner["research_scope"])

    heading("Literature Review")
    if literature.get("low_coverage"):
        doc.add_paragraph("⚠ Low coverage: fewer than 3 relevant papers were retrieved.")
    similar_work = literature.get("similar_work")
    if similar_work:
        pct = round(similar_work.get("similarity", 0) * 100)
        p = doc.add_paragraph()
        p.add_run(f"⚠ Possible existing work: {pct}% similar to ")
        _add_hyperlink(p, similar_work.get("paper_url", ""), similar_work.get("paper_title", ""))
        p.add_run(" — check it isn't already covered.")
    for paper in literature.get("papers", []):
        authors = ", ".join(paper.get("authors", []))
        p = doc.add_paragraph(style="List Bullet")
        _add_hyperlink(p, paper.get("url", ""), paper.get("title", "Untitled"), bold=True)
        p.add_run(f" ({paper.get('year', 'n.d.')}) — {authors}")
        if paper.get("url_verified") is False:
            p.add_run("  [link unreachable]").italic = True
        if paper.get("summary"):
            doc.add_paragraph(paper["summary"]).paragraph_format.left_indent = Pt(18)
        if paper.get("citation_apa"):
            cite_p = doc.add_paragraph()
            cite_p.paragraph_format.left_indent = Pt(18)
            cite_p.add_run(f"Citation (APA): {paper['citation_apa']}").italic = True
    if not literature.get("papers"):
        doc.add_paragraph("No papers could be retrieved for this topic.")

    heading("Research Gaps")
    for t in gap.get("current_trends", []):
        bullet(f"Trend: {t}")
    for g in gap.get("gaps", []):
        based_on = ", ".join(g.get("based_on_papers", [])) or "—"
        p = doc.add_paragraph(style="List Bullet")
        p.add_run("Gap: ").bold = True
        p.add_run(g.get("gap_description", ""))
        detail = doc.add_paragraph(f"Based on: {based_on}; {g.get('confidence_note', '')}")
        detail.paragraph_format.left_indent = Pt(18)
        detail.runs[0].italic = True
    for d in gap.get("future_directions", []):
        bullet(f"Future direction: {d}")

    heading("Methodology")
    doc.add_paragraph("Datasets:").runs[0].bold = True
    for d in methodology.get("datasets", []):
        badge = "verified" if d.get("verified") else "unverified"
        if d.get("url_verified") is False:
            badge += ", link unreachable"
        link = d.get("kaggle_url") or d.get("url") or ""
        p = doc.add_paragraph(style="List Bullet")
        _add_hyperlink(p, link, d.get("name", ""))
        p.add_run(f" — {badge}")

    doc.add_paragraph("Algorithms:").runs[0].bold = True
    for a in methodology.get("algorithms", []):
        badge = "verified" if a.get("verified") else "unverified"
        if a.get("url_verified") is False:
            badge += ", link unreachable"
        p = doc.add_paragraph(style="List Bullet")
        if a.get("url"):
            _add_hyperlink(p, a["url"], a.get("name", ""))
        else:
            p.add_run(a.get("name", ""))
        p.add_run(f" — {badge}")

    bold_line("Evaluation metrics", ", ".join(methodology.get("evaluation_metrics", [])))
    doc.add_paragraph("Implementation steps:").runs[0].bold = True
    for step in methodology.get("implementation_steps", []):
        doc.add_paragraph(step, style="List Number")
    if methodology.get("unverified_suggestions"):
        doc.add_paragraph("Unverified suggestions (not checked against a real source):").runs[0].bold = True
        for s in methodology["unverified_suggestions"]:
            bullet(s)

    heading("Reviewer Feedback")
    bold_line(
        "Scores",
        f"Novelty {reviewer.get('novelty_score', 'n/a')} · Quality {reviewer.get('quality_score', 'n/a')}",
    )
    doc.add_paragraph("Rule checks:").runs[0].bold = True
    for check, passed in reviewer.get("rule_check_results", {}).items():
        bullet(f"{'PASS' if passed else 'FAIL'} — {check.replace('_', ' ')}")
    if reviewer.get("strengths"):
        doc.add_paragraph("Strengths:").runs[0].bold = True
        for s in reviewer["strengths"]:
            bullet(s)
    if reviewer.get("weaknesses"):
        doc.add_paragraph("Weaknesses:").runs[0].bold = True
        for w in reviewer["weaknesses"]:
            bullet(w)
    if reviewer.get("missing_components"):
        doc.add_paragraph("Missing components:").runs[0].bold = True
        for m in reviewer["missing_components"]:
            bullet(m)

    footer = _flags_footer(session)
    if footer:
        doc.add_heading("Caveats", level=1)
        for f in footer:
            bullet(f"⚠ {f}")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
